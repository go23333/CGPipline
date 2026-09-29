#coding=utf-8

import os
import openpyxl as op


from PySide2.QtWidgets import *
from PySide2.QtGui import *
from PySide2.QtCore import *

from dayu_widgets.message import MMessage
from dayu_widgets.label import MLabel
from dayu_widgets.line_edit import MClickBrowserFolderToolButton,MLineEdit
from dayu_widgets.push_button import MPushButton




class mw(QWidget):



    def __init__(self, parent=None):
        super(mw,self).__init__(parent)

        
        self.uii()

    def uii(self):   
        self.setWindowTitle(u'动画fbx完整性检查工具')
        self.resize(350,150)
        lay=QVBoxLayout()

        folder_lay=QVBoxLayout()

        #导入文件
        self.excel_path=MLineEdit().file(filters=['xlsx']).medium()
        self.excel_path.setMinimumHeight(30)
        self.excel_path.setPlaceholderText(u"选择需要对比信息的Excel文件")
        self.fbx_path=MLineEdit().folder().medium()
        self.fbx_path.setMinimumHeight(30)
        self.fbx_path.setPlaceholderText(u"选择需要对比信息的fbx文件夹")
        create_folder=MPushButton(text=u"执行")
        create_folder.clicked.connect(self.excute)
        

        folder_lay.addWidget(self.excel_path)
        folder_lay.addWidget(self.fbx_path)
        folder_lay.addWidget(create_folder)
        

        lay.addLayout(folder_lay)
        self.setLayout(lay)



    def excute(self):
        excel_path=self.excel_path.text()
        fbx_path=self.fbx_path.text()
        df=op.load_workbook(excel_path,data_only=True)
        sheet=df.get_sheet_by_name('Sheet1')
        
        
        all_list=[]
        excel_dict={}
        i=0
        #确定行和列
        rowcount=sheet.max_row
        colcount=3
        #读取所需的Excel内容
        for i in range(1,rowcount+1):
            list=[]
            for j in range(1,colcount+1):
                list.append(sheet.cell(row=i,column=j).value)
            if list[2]:
                for i in range(list[2]):
                    #all_list.append(list[1])
                    if '_BG' not in list[1]:
                        try:
                            excel_dict[list[0]].append(list[1].split('.')[0])
                        except:
                            excel_dict[list[0]]=[]
                            excel_dict[list[0]].append(list[1].split('.')[0])
        #print(excel_dict)
        #print('')
        
        fbx_dict={}
        for dirpath, dirnames, filenames in os.walk(fbx_path):
            for filename in filenames:
                fbx_list=[]
                if '_CH' in filename or '_Pro' in filename:
                    if '_CH' in filename:
                        suffix='_CH'
                        fbx_list=filename.split('_CH')[0].split('_an_')
                    elif '_Pro' in filename:
                        suffix='_Pro'
                        fbx_list=filename.split('_Pro')[0].split('_an_')
                    try:
                        fbx_dict[fbx_list[0]].append(fbx_list[1]+suffix)
                    except:
                        fbx_dict[fbx_list[0]]=[]
                        fbx_dict[fbx_list[0]].append(fbx_list[1]+suffix)
                    
        
        
        keys1=set(excel_dict.keys())
        keys2=set(fbx_dict.keys())
        diff_set=keys1-keys2
        equal_set=keys1 & keys2
        equal_list=[key for key in equal_set]
        diff_list=[key for key in diff_set]
        equal_list.sort()
        diff_list.sort()
        
        
        result_list=[]
        for k in diff_list:
            result_list.append([k,'all'])
        
        
        for key in equal_list:
            list1=excel_dict[key]
            new_equal_list=[x for x in excel_dict[key]]
            list2=fbx_dict[key]
            #diff_set=set(list1)-set(list2)
            #new_equal_list=[]
            #for k in diff_set:
            #    new_equal_list.append(k)
            
            #print(list1)
            #print(list2)
            
            index_list=[]
            i=0
            for v1 in list1:
                for v2 in list2:
                    if v2 == v1:
                        #new_equal_list.append()
                        index_list.append(i)
                        list2.remove(v2)
                        break
                i+=1
            for index in index_list:
                v=list1[index]
                new_equal_list.remove(v)
            if new_equal_list:
                result_list.append([key,new_equal_list])
        
        print('')
        for value in result_list:
            print(value)




print('')


def showUI():
    app=QApplication.instance()
    global win
    win=mw()
    win.show()
    win.setWindowFlags(Qt.WindowMinimizeButtonHint | Qt.WindowCloseButtonHint | Qt.WindowStaysOnTopHint)
    win.show()
    app.exec_()




if __name__=='__main__':
    showUI()
