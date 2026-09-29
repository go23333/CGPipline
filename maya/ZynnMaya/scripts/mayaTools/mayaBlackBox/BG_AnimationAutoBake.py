# #coding=utf-8

# import sys
# scriptsPath = r'Y:\scripts\mayaTools\site-packages'
# if scriptsPath not in sys.path:
#     sys.path.insert(0, scriptsPath)

from maya import cmds,mel
import time
import os
import tempfile
import getpass
import string
import openpyxl as xls

import json

print('AAB1.3')



camExAttrs = [u"translateX",u"translateY",u"translateZ",u"rotateX",u"rotateY",u"rotateZ","scaleX",u"scaleY",u"scaleZ"]
def findSceneName():
    ScName=cmds.file(q=1,ns=1)
    return ScName
def cameraCheckAov():
    #sheetRow = self.sheetRow()
    scName = findSceneName()
    #rowNum=sheetRow+1
    #print rowNum
    cameraAll=cmds.ls(type='camera')
    fileN = findSceneName()
    fileNames=fileN.split('_')
    islockApp=[]
    exAttr = camExAttrs
    cameraApp=[]
    leftParent=''
    for c in cameraAll:
        if 'left' in c:
            leftParent = cmds.listRelatives(c,p=1)
            cmds.delete(leftParent)
    cameraAll=cmds.ls(type='camera')
    for cc in cameraAll:
        if 'lookThrough_faceCtrlCameraShape' not in cc and 'faceCameraShape' not in cc:
            cmds.setAttr('%s.renderable'%cc,0)
            #cmds.setAttr('Ep001_sc009_004_001_020_cam.renderable',1)
            cameraApp.append(cc)
    # if len(cameraApp)>5:
    #     cmds.error(u'\u76f8\u673a\u5b58\u5728\u4e24\u4e2a\u4ee5\u4e0a----------------------------\u8f6c\u6362\u5931\u8d25----------------------------')
    #     #self.blankRowInsert(rowNum,1,scName,u'\u76f8\u673a\u5b58\u5728\u4e24\u4e2a\u4ee5\u4e0a')
    # elif  len(cameraApp)==5:
        #for b in cameraAll:
            #cmds.setAttr('%s.renderable'%b,0)
    for c in cameraApp:
        if '_' in c:
            cameraSan=c.split('_')
            if cameraSan[1] == fileNames[1] and cameraSan[2] == fileNames[2]:
                cameraT=cmds.listRelatives( c, p=True )                        
                cmds.playbackOptions(min = cameraSan[3])
                cmds.playbackOptions(max = cameraSan[4])
                cmds.playbackOptions(ast = cameraSan[3])
                cmds.playbackOptions(aet = cameraSan[4])
                for ce in exAttr:
                    if cmds.objExists(cameraT[0]+'.'+ce):
                        islock=cmds.getAttr(cameraT[0]+'.'+ce,l=1)
                        if islock==False:
                            islockApp.append(islock)
                return c
            else:
                cmds.error(u'\u76f8\u673a\u547d\u540d\u4e0d\u5bf9----------------------------\u8f6c\u6362\u5931\u8d25----------------------------')
                #self.blankRowInsert(rowNum,1,scName,u'\u76f8\u673a\u547d\u540d\u4e0d\u5bf9')

def TwoCameraCheckAov():
    camTwoPoint = ''
    cameraTwoPoint=cmds.ls(type='lookAt')
    fileNames=cmds.file(q=1,sn=1,shn=1).split('_')
    if len(cameraTwoPoint)==1:
        cameraName = cmds.listRelatives(cameraTwoPoint,c=1)
        for c in cameraName:
            cam = cmds.listRelatives(c,c=1)
            if cmds.nodeType(cam) == 'camera':
                camTwoPoint = cam[0]
                cameraSan=camTwoPoint.split('_')
                if cameraSan[1] == fileNames[1] and cameraSan[2] == fileNames[2]:
                    cmds.playbackOptions(min = cameraSan[3])
                    cmds.playbackOptions(max = cameraSan[4])
                    cmds.playbackOptions(ast = cameraSan[3])
                    cmds.playbackOptions(aet = cameraSan[4])                      
                    return camTwoPoint
                else:
                    cmds.error(u'\u76f8\u673a\u547d\u540d\u4e0d\u5bf9----------------------------\u8f6c\u6362\u5931\u8d25----------------------------')
            #return camTwoPoint
    elif len(cameraTwoPoint)>1:
        cmds.error(u'\u4e24\u70b9\u76f8\u673a\u4e0d\u5bf9')

def ConvertNewCam(start_offset,end_offset):
    oldCam = cameraCheckAov()
    TwoPointCam = TwoCameraCheckAov()
    if oldCam and TwoPointCam is None:
        oldCamTx = cmds.listRelatives(oldCam,p=1)[0]
        cameraSplit = oldCam.split('_')
        cmds.currentTime(cameraSplit[3])
        TX = cmds.getAttr('%s.tx'%oldCamTx)
        TY = cmds.getAttr('%s.ty'%oldCamTx)
        TZ = cmds.getAttr('%s.tz'%oldCamTx)
        RX = cmds.getAttr('%s.rx'%oldCamTx)
        RY = cmds.getAttr('%s.ry'%oldCamTx)
        RZ = cmds.getAttr('%s.rz'%oldCamTx)
        newCam = cmds.camera()
        cmds.setAttr('%s.tx'%newCam[0],TX)
        cmds.setAttr('%s.ty'%newCam[0],TY)
        cmds.setAttr('%s.tz'%newCam[0],TZ)
        cmds.setAttr('%s.rx'%newCam[0],RX)
        cmds.setAttr('%s.ry'%newCam[0],RY)
        cmds.setAttr('%s.rz'%newCam[0],RZ)
        
        minClip = cmds.getAttr('%s.nearClipPlane'%oldCam)
        farClip = cmds.getAttr('%s.farClipPlane'%oldCam)
        horizontalFilm = cmds.getAttr('%s.horizontalFilmAperture'%oldCam)
        verticalFilm = cmds.getAttr('%s.verticalFilmAperture'%oldCam)
        filmFit = cmds.getAttr('%s.filmFit'%oldCam)
        focalLengthOld = cmds.getAttr('%s.focalLength'%oldCam)
        
        #cmds.setAttr('%s.focalLength'%newCam[1],focalLength)
        cmds.setAttr('%s.nearClipPlane'%newCam[1],minClip)
        cmds.setAttr('%s.farClipPlane'%newCam[1],farClip)
        cmds.setAttr('%s.horizontalFilmAperture'%newCam[1],horizontalFilm)
        cmds.setAttr('%s.verticalFilmAperture'%newCam[1],verticalFilm)
        cmds.setAttr('%s.filmFit'%newCam[1],filmFit)
        cmds.parentConstraint(oldCamTx,newCam[0],w=1,mo=0,st = 'none',sr = 'none')

        minTime = int(cameraSplit[3])
        maxTime = int(cameraSplit[4])+1
        cmds.bakeResults(newCam[0],simulation=True,sampleBy=1,hierarchy="below",t=(cameraSplit[3],cameraSplit[4]),at=['tx','ty','tz','rx','ry','rz','sx','sy','sz'])
            
        for ii in range(minTime,maxTime):
            cmds.currentTime(ii)
            focalLength = cmds.getAttr('%s.focalLength'%oldCam)
            cmds.setAttr('%s.focalLength'%newCam[1],focalLength)
            cmds.setKeyframe(newCam[1],at = ['focalLength'],t=[ii,ii])
            #cmds.bakeResults(newCam[1],simulation=True,sampleBy=1,hierarchy="below",t=(cameraSplit[3],cameraSplit[4]),at=['focalLength'])
        
        cmds.currentTime(cameraSplit[3])
        focalLengthOne = cmds.getAttr('%s.focalLength'%oldCam)
        
        for j in range(minTime-start_offset,minTime):
            cmds.currentTime(j)
            cmds.setAttr('%s.tx'%newCam[0],TX)
            cmds.setAttr('%s.ty'%newCam[0],TY)
            cmds.setAttr('%s.tz'%newCam[0],TZ)
            cmds.setAttr('%s.rx'%newCam[0],RX)
            cmds.setAttr('%s.ry'%newCam[0],RY)
            cmds.setAttr('%s.rz'%newCam[0],RZ)
            cmds.setAttr('%s.nearClipPlane'%newCam[1],minClip)
            cmds.setAttr('%s.farClipPlane'%newCam[1],farClip)
            cmds.setAttr('%s.horizontalFilmAperture'%newCam[1],horizontalFilm)
            cmds.setAttr('%s.verticalFilmAperture'%newCam[1],verticalFilm)
            cmds.setAttr('%s.filmFit'%newCam[1],filmFit)
            cmds.setAttr('%s.focalLength'%newCam[1],focalLengthOne)
            cmds.setKeyframe(newCam[1],at = ['focalLength'],t=[j,j])
            cmds.setKeyframe(newCam[0],at = ['tx','ty','tz','rx','ry','rz','sx','sy','sz'],t=[j,j])

        cmds.currentTime(cameraSplit[4])
        TX = cmds.getAttr('%s.tx'%oldCamTx)
        TY = cmds.getAttr('%s.ty'%oldCamTx)
        TZ = cmds.getAttr('%s.tz'%oldCamTx)
        RX = cmds.getAttr('%s.rx'%oldCamTx)
        RY = cmds.getAttr('%s.ry'%oldCamTx)
        RZ = cmds.getAttr('%s.rz'%oldCamTx)
        focalLengthLast = cmds.getAttr('%s.focalLength'%oldCam)
        minClip = cmds.getAttr('%s.nearClipPlane'%oldCam)
        farClip = cmds.getAttr('%s.farClipPlane'%oldCam)
        horizontalFilm = cmds.getAttr('%s.horizontalFilmAperture'%oldCam)
        verticalFilm = cmds.getAttr('%s.verticalFilmAperture'%oldCam)
        filmFit = cmds.getAttr('%s.filmFit'%oldCam)
        
        for t in range(maxTime,maxTime+end_offset):
            cmds.currentTime(t)
            cmds.setAttr('%s.tx'%newCam[0],TX)
            cmds.setAttr('%s.ty'%newCam[0],TY)
            cmds.setAttr('%s.tz'%newCam[0],TZ)
            cmds.setAttr('%s.rx'%newCam[0],RX)
            cmds.setAttr('%s.ry'%newCam[0],RY)
            cmds.setAttr('%s.rz'%newCam[0],RZ)
            cmds.setAttr('%s.nearClipPlane'%newCam[1],minClip)
            cmds.setAttr('%s.farClipPlane'%newCam[1],farClip)
            cmds.setAttr('%s.horizontalFilmAperture'%newCam[1],horizontalFilm)
            cmds.setAttr('%s.verticalFilmAperture'%newCam[1],verticalFilm)
            cmds.setAttr('%s.filmFit'%newCam[1],filmFit)
            cmds.setAttr('%s.focalLength'%newCam[1],focalLengthLast)
            cmds.setKeyframe(newCam[1],at = ['focalLength'],t=[t,t])
            cmds.setKeyframe(newCam[0],at = ['tx','ty','tz','rx','ry','rz','sx','sy','sz'],t=[t,t])

  
        Camchild = cmds.listRelatives(newCam[0],c=1)
        for c in Camchild:
            if cmds.nodeType(c)=='parentConstraint':
                cmds.delete(c)
        #move right 10
        #oldCamTx = 'Ep002_sc006_002_001_068_cam'
        oldCamSp = oldCamTx.split('_')
        #加str()防打包pyd失败
        oldCamTxNew = str('%s_%s_%s_%s_%03d_cam'%(oldCamSp[0],oldCamSp[1],oldCamSp[2],oldCamSp[3],int(oldCamSp[4])+start_offset+end_offset))
        print(newCam[0],oldCamTxNew+('_export'))
        newName = cmds.rename(newCam[0],oldCamTxNew+('_export'))
        cmds.copyKey(newName, time=(minTime-start_offset,maxTime+end_offset))
        cmds.pasteKey(newName, time=(minTime,minTime),option="replaceCompletely")
        
        return newName
    elif TwoPointCam:
        TwoPointCamT = cmds.listRelatives(TwoPointCam,p=1)[0]
        TwoPointCamTSp = TwoPointCamT.split('_')
        cmds.currentTime(TwoPointCamTSp[3])
        newTwoCam = cmds.camera()
        minTime = string.atoi(TwoPointCamTSp[3])-start_offset
        maxTime = string.atoi(TwoPointCamTSp[4])+end_offset
        
        minClip = cmds.getAttr('%s.nearClipPlane'%TwoPointCamT)
        farClip = cmds.getAttr('%s.farClipPlane'%TwoPointCamT)
        horizontalFilm = cmds.getAttr('%s.horizontalFilmAperture'%TwoPointCamT)
        verticalFilm = cmds.getAttr('%s.verticalFilmAperture'%TwoPointCamT)
        filmFit = cmds.getAttr('%s.filmFit'%TwoPointCamT)
        
        cmds.setAttr('%s.nearClipPlane'%newTwoCam[1],minClip)
        cmds.setAttr('%s.farClipPlane'%newTwoCam[1],farClip)
        cmds.setAttr('%s.horizontalFilmAperture'%newTwoCam[1],horizontalFilm)
        cmds.setAttr('%s.verticalFilmAperture'%newTwoCam[1],verticalFilm)
        cmds.setAttr('%s.filmFit'%newTwoCam[1],filmFit)
        newName = cmds.rename(newTwoCam[0],TwoPointCamT+('_export'))
        newNameShape = cmds.listRelatives(newName,c=1)[0]
        for i in range(minTime,maxTime):
            cmds.currentTime(i)
            pos = cmds.xform(TwoPointCamT,q=1,ws=1,t=1)
            rot = cmds.xform(TwoPointCamT,q=1,ws=1,ro=1)
            scal = cmds.xform(TwoPointCamT,q=1,ws=1,s=1)
            cmds.setAttr(newName+'.tx',pos[0])
            cmds.setAttr(newName+'.ty',pos[1])
            cmds.setAttr(newName+'.tz',pos[2])
            cmds.setAttr(newName+'.rx',rot[0])
            cmds.setAttr(newName+'.ry',rot[1])
            cmds.setAttr(newName+'.rz',rot[2])
            cmds.setAttr(newName+'.sx',scal[0])
            cmds.setAttr(newName+'.sy',scal[1])
            cmds.setAttr(newName+'.sz',scal[2])
            focalLength = cmds.getAttr('%s.focalLength'%TwoPointCamT)
            cmds.setAttr('%s.focalLength'%newNameShape,focalLength)
            cmds.setKeyframe(newName,at = ['tx','ty','tz','rx','ry','rz'],t=[i,i])
            cmds.setKeyframe(newNameShape,at = ['focalLength'],t=[i,i])
        return TwoPointCamT
    else:
        cmds.error(u'\u76f8\u673a\u547d\u540d\u4e0d\u89c4\u8303')
        return
    
def saveCam(offset_franme,export_path):
    #获取偏移值
    start_offset=int(offset_franme[0])
    end_offset=int(offset_franme[1])

    cameraName = ConvertNewCam(start_offset=start_offset,end_offset=end_offset)
    camPath = export_path+'/CamPath/'
    if os.path.exists(camPath)==False:
        os.makedirs(camPath)
    newcamSP = cameraName.split('_')
    cameraName1 = '%s_%s_%s_cam'%(newcamSP[0],newcamSP[1],newcamSP[2])
    cmds.file('%s%s'%(camPath,cameraName1),force=True, exportSelected=True, options = "v=0;p=17;f=0",typ = "FBX export" ,pr=1)
    #cmds.delete(cameraName)
    # print (u'\u76f8\u673a\u5bfc\u51fa\u5230%s\u6587\u4ef6\u5939'%camPath)




class RefPrintXml():
    camExAttrs = [u"translateX",u"translateY",u"translateZ",u"rotateX",u"rotateY",u"rotateZ","scaleX",u"scaleY",u"scaleZ"]
    
    sheetNum = 'Sheet1'
    CurrentTime=time.strftime('%Y_%m_%d',time.localtime(time.time()-2592000))
    def __init__(self,export_path,excel_base_path):
        self.export_path = export_path
        self.excel_base_path = excel_base_path
        self.XMLfileName = "Z:\\NScripts_2016\\NScripts\\noRemove\\Null.xlsx"
        if not os.path.isfile(self.XMLfileName):
            self.XMLfileName = self.excel_base_path

        pass
        
    def sheetRow(self,P):
        if os.path.exists(P):
            fn=P
        else:
            fn=self.XMLfileName
        sn=self.sheetNum
        wb = xls.load_workbook(fn)
        sheet = wb.get_sheet_by_name(sn)
        return len(sheet.rows)    
            
    def blankRowInsert(self,N, M,P,A1=None,B1=None,C1=None,D1=None):
        if os.path.exists(P):
            fn=P
        else:
            fn=self.XMLfileName
        sp=P
        ct=self.CurrentTime
        sn=self.sheetNum
        wb = xls.load_workbook(fn)
        sheet = wb.get_sheet_by_name(sn)
        myList = self.dumpDataToList(sheet,P)
        self.insertLine(myList, N, M, sheet.max_column,A1,B1,C1,D1)
        self.setValue(sheet,myList)
        wb.save(P)
        #wb.close()
        return len(myList)
     
    def dumpDataToList(self,sheet,P):
        if os.path.exists(P):
            fn=P
        else:
            fn=self.XMLfileName
        sp=P
        ct=self.CurrentTime
        sn=self.sheetNum
        wb = xls.load_workbook(fn)
        sn=self.sheetNum
        sheet = wb.get_sheet_by_name(sn)
        listResult = []
        for i in range(1,sheet.max_row + 1):
            lineData = []
            for j in range(1,sheet.max_column +1):
                cell = sheet.cell(row = i, column = j)
                lineData.append(cell.value)
            listResult.append(lineData)
        return listResult
         
    def insertLine(self,aList, N , M, maxColumn,A1=None,B1=None,C1=None,D1=None):
        for _ in range(1,M + 1):
            aList.insert(N, [A1,B1,C1,D1] )
            #aList.insert(N, [] * maxColumn)
            #aList.insert(N, ['ccc'] * maxColumn)
    
    def setValue(self,sheet,list):
        for i in range(1, len(list) + 1):
            for j in range(1, len(list[i-1]) + 1):
                cell = sheet.cell(row=i, column=j)
                cell.value = list[i-1][j-1]
                sheet.column_dimensions['A'].width=20
        for i in range(1, len(list) + 1):
            for j in range(1, len(list[i-1]) + 1):
                cell = sheet.cell(row=i, column=j)
                cell.value = list[i-1][j-1]
                sheet.column_dimensions['B'].width=50
                #sheet.column_dimensions['B'].width=50
    def findSceneName(self):
        ScName=cmds.file(q=1,ns=1)
        return ScName
 
    def getFiles(self,path,suffix):
        return [os.path.join(root, file) for root, dirs, files in os.walk(path) for file in files if file.endswith(suffix)]
            
    def getProjNa(self):
        projectName = cmds.file(q=1,sn=1).split('/')
        projectList = os.listdir('Y:')
        if projectName and projectList:
            for a in projectName:
                for b in projectList:
                    if a == b :
                        return a
        else:
            cmds.error(u'\u8bf7\u786e\u4fdd\u5f53\u524d\u6587\u4ef6\u662f\u6253\u5f00\u7684\u5e76\u4e14Y\u76d8\u4e0b\u6709\u9879\u76ee\u6587\u4ef6\u5939')
            return 0
            
    def ImportFideScenceName(self):
        fileName=cmds.file(q=1,sn=1,shn=1)
        fileSp=fileName.split("_")
        return fileName
    
    def ImportABCFile(self,dic):
        fileName=self.ImportFideScenceName()
        projNa = self.getProjNa()
        if not projNa:
            projNa = 'Temp'
        fileName1 = cmds.file(q=1,ns=1)
        fileNameSplit = fileName1.split('_')
        # outputMovPath = export_path+'/RefPrint'
        # outputMovPath = 'E:/RefPrint'
        # if not os.path.exists(outputMovPath):
        #     outputMovPath = 'D:/RefPrint'
        outputMovPath = self.export_path+'/RefPrint'
        print("output path:", outputMovPath)
        outputXmlPathAll = ("%s/%s/%s/%s"%(outputMovPath,projNa.lower(),fileNameSplit[0],fileNameSplit[1]))
        #print outputXmlPathAll
        if not os.path.exists(outputXmlPathAll):
            os.makedirs(outputXmlPathAll)
        saveXmlPath='%s/%s_%s_%s.xlsx'%(outputXmlPathAll,fileNameSplit[0],fileNameSplit[1],self.CurrentTime)

        sheetRow = self.sheetRow(saveXmlPath)
        rowNum=sheetRow+1
        shotName = '%s_%s_%s'%(fileNameSplit[0],fileNameSplit[1],fileNameSplit[2])
        self.blankRowInsert(rowNum,1,saveXmlPath,None,shotName)
        #ref = cmds.file(q=1,r=1)
        rowNum += 1
        for k,v in dic.items():
            try:
                self.blankRowInsert(rowNum,1,saveXmlPath,shotName,k,v)
                #self.blankRowInsert(rowNum,2,saveXmlPath,v)
            except WindowsError:
                pass

    def ref_check(self):
        refile = cmds.file(q=1, r=1)
        refile_c = []
        count_dict = {}
        for r in refile:
            fileName = r.split('/')[-1]
            if '{' in fileName:
                fileName = fileName.split('{')[0]
            refile_c.append(fileName)

        for item in refile_c:
            if item in count_dict:
                count_dict[item] += 1
            else:
                count_dict[item] = 1  
        self.ImportABCFile(count_dict)     
        #return count_dict
   
    def refChaFind(self):
        refile = cmds.file(q=1, r=1)
        refNodeChaNameAll = []
        refNodeName = []
        for r in refile:
            # refileSplit = r.split('/')
            if r.split('/')[4]=='Character':
                try:
                    refNode = cmds.referenceQuery(r, ns=True)
                    refNodeName = refNode.split(':')[-1] + ':Group'
                    refNodeChaNameAll.append(refNodeName)
                except RuntimeError:
                    pass
        return refNodeChaNameAll

def openJson():
    temp_path = tempfile.gettempdir()
    # temp_path=temp_path.replace('\\','/')
    json_path = os.path.join(temp_path,'mayajson.json')
    # json_path = 'D:/Documents/maya/scripts/mayajson.json'
    if os.path.exists(json_path):
        with open(json_path,'r') as js_file:
            json_data=json.load(js_file)
        if json_data :
            return json_data
        
        
def txtWrite(data):
    desktop_log_name ="maya_AnExport_log.txt"
    try:
        desktop_path = "d:/Desktop/"
        if not os.path.exists(desktop_path):
            # desktop_path = 'C:/Users/Administrator/Desktop/'
            # username = getpass.getuser()
            # desktop_path = desktop_path.replace('Administrator',username)
            desktop_path = os.path.join(os.path.join(os.environ['USERPROFILE']), 'Desktop')

    except:
        desktop_path = 'Y:/temp/'
        if not os.path.exists(desktop_path):
            os.makedirs(desktop_path)
    
        
    desktop_log_path = desktop_path+desktop_log_name
    file_path = desktop_log_path
    with open(file_path, "a") as file:
        try:
            file.write(data+"\n")
        except:
            file.write("data write error\n")


def animationLayerDetection(fbx_export_path,maya_path):

    log_path = fbx_export_path+'/an_layer_error.txt'
    #获取动画层
    animLayers = cmds.ls(type='animLayer')
    for animLayer in animLayers:
        if animLayer != 'BaseAnimation':
            data = 'redundant animation layers     '+maya_path
            with open(log_path, "a") as file:
                file.write(data+"\n")
            break

    


# def start():
#     maya_dict=openJson()
#     maya_path=maya_dict['maya_path']
#     fbx_export_path=maya_dict['fbx_export_path']
#     start_endK_d=maya_dict['frame']
#     offset_franme=maya_dict['offset_frame']
#     start_frame=int(start_endK_d['frame'][0])#获取关键帧起始信息
#     end_frame=int(start_endK_d['frame'][1])
#     print(start_frame,end_frame)




def execute():
    maya_dict=openJson()
    maya_path=maya_dict['maya_path']
    fbx_export_path=maya_dict['fbx_export_path']
    start_endK_d=maya_dict['frame']
    offset_franme=maya_dict['offset_frame']
    root_path=maya_dict['root_path']
            

    #print(cmds.file(q=1,sn=1))
    cmds.loadPlugin('fbxmaya')
    try:
        cmds.loadPlugin('ZYNNnode')
    except:
        pass
    cmds.file(modified=0)
    # cmds.file(maya_path,o=1,executeScriptNodes=False,ignoreVersion=True)
    mel.eval('file -lar -f -options "v=0;" -esn false -ignoreVersion -o \"%s\";'%(maya_path))

    import sys
    scriptsPath = root_path+'/site-packages'
    if scriptsPath not in sys.path:
        sys.path.insert(0, scriptsPath)

    animationLayerDetection(fbx_export_path,maya_path)
    #为日志写入文件信息
    timestamp = time.time()
    current_time = time.ctime(timestamp)
    content = '###########################\n'+maya_path+'    '+current_time
    txtWrite(content)

    #创建引用表
    #"Z:\\NScripts_2016\\NScripts\\noRemove\\Null.xlsx"
    excel_base_path = root_path+'/other/Null.xlsx'
    RefPrintXml(fbx_export_path,excel_base_path).ref_check()
    #导出摄像机
    try:
        saveCam(offset_franme,fbx_export_path)
    except:
        content = maya_path+'    camera export error'
        txtWrite(content)

    time.sleep(5)
    #将帧速率设置为25
    cmds.currentUnit( time='pal' )
    

    start_frame=start_endK_d[0]#获取关键帧起始信息
    end_frame=start_endK_d[1]
    
    #获取偏移值
    start_offset=int(offset_franme[0])
    end_offset=int(offset_franme[1])
    
    start_key=start_frame
    end_key=end_frame+start_offset+end_offset
    
    #设置时间轴范围
    cmds.playbackOptions( minTime=start_key, maxTime=end_key )

    #获取全部名称空间名
    space_names=cmds.namespaceInfo( lon=True )
    new_space_names=[]
    for space_name in space_names:      #为名称空间调整排序,道具在前,角色在后
        if cmds.objExists(space_name+':Root_M'):
            new_space_names.append(space_name)
        elif cmds.objExists(space_name+':*_GuGe_G'):
            new_space_names.insert(0,space_name)

    print(space_names)
    file_names=[]
    for space_name in new_space_names:
        print('')
        if cmds.objExists(space_name+':Root_M'):
            #选择所有对应名称空间的对象
            
            print('ch',space_name)
            cmds.select(space_name+':DeformationSystem')
            
            try:
                #cmds.select(space_name+':Geometry',add=1)
                select_obj,file_name=bake(start_offset,end_offset,start_frame,end_frame)
                #多个同名文件添加数字后缀
                i=1
                new_file_name=file_name
                while new_file_name in file_names:
                    new_file_name=file_name+str(i)
                    i+=1
                file_names.append(new_file_name)
                export_path = exportFBX(fbx_export_path,new_file_name,select_obj,start_key,end_key)
                print('ch',new_file_name)
                txtWrite('  ch '+new_file_name+'    '+export_path)
            except:
                print('!!!!!!!!!!!!'+space_name+' export fail')
                txtWrite('  !!!!!!!!!!!!'+space_name+' export fail')
            
        elif cmds.objExists(space_name+':*_GuGe_G'):

            print('pro',space_name)
            cmds.select(space_name+':*_GuGe_G')
            
            try:
                # cmds.select(space_name+':*_Pro_Mo',add=1)
                select_obj,file_name=bake(start_offset,end_offset,start_frame,end_frame)
                #多个同名文件添加数字后缀
                i=1
                new_file_name=file_name
                while new_file_name in file_names:
                    new_file_name=file_name+str(i)
                    i+=1
                file_names.append(new_file_name)
                export_path = exportFBX(fbx_export_path,new_file_name,select_obj,start_key,end_key)
                print('pro',new_file_name)
                txtWrite('  pro '+new_file_name+'    '+export_path)
            except:
                print('!!!!!!!!!!!!'+space_name)
                txtWrite('  !!!!!!!!!!!!'+space_name+' export fail')
                


def bake(start_offset,end_offset,start_frame,end_frame):
    
    obj=cmds.ls(sl=1)

    #导入选定对象引用并删除选定对象的名称空间
    path_reference=cmds.referenceQuery(obj[0], f=True )
    cmds.file(path_reference, ir=1)
    # cmds.parent(obj,world=True)#解除组
    #获取对象空间命名
    obj_namespace=str(obj[0]).split(':',1)[0]+':'
    #通过引用路径获取文件名
    file_name = path_reference.split('/')[-1].split('.')[0]
    cmds.select(obj)
    obj_new=cmds.ls(sl=1)
    
    #选择需要烘焙的对象
    cmds.select(obj_new,hi=1)
    
    #烘焙
    obj_bake_l=[]
    obj_list=cmds.ls(sl=1)
    i=0
    while i<len(obj_list):
        obj_bake_l.append(str(obj_list[i]))
        i+=1
    obj_bake_a=str(obj_bake_l)[2:-2]
    obj_bake=obj_bake_a.replace("'",'"')
    
    #设置烘焙起始结束帧
    print(start_frame,end_frame)
    mel.eval('''bakeResults -simulation true -t "%s:%s" -sampleBy 1 -oversamplingRate 1 -disableImplicitControl true -preserveOutsideKeys true -sparseAnimCurveBake false -removeBakedAttributeFromLayer false -removeBakedAnimFromLayer false -bakeOnOverrideLayer false -minimizeRotation true -controlPoints false -shape true {"%s"}'''%(start_frame,end_frame,obj_bake))
    
    
    #选择所有烘焙的对象
    cmds.select(obj_new,hi=1)
    #将所有对象放入列表
    frame_list=cmds.ls(sl=1)
    #偏移关键帧
    cmds.keyframe( frame_list,relative=True,timeChange=start_offset)
    
    #烘焙前后帧
    if start_offset!=0:
        cmds.bakeResults( frame_list, t=(start_frame,start_frame+start_offset+1), simulation=True,preserveOutsideKeys=True )
    if end_offset!=0:
        cmds.bakeResults( frame_list, t=(end_frame,end_frame+start_offset+end_offset), simulation=True,preserveOutsideKeys=True )


    #删除其他兄弟组(避免删除空间命名失败)
    cmds.select(obj_new)
    target_group=cmds.ls(sl=1)[0]
    parent_group = cmds.listRelatives(target_group, parent=True, fullPath=True)
    all_children = cmds.listRelatives(parent_group, children=True, fullPath=True)
    cmds.select(clear=1)
    for children in all_children:
        if target_group !=children.split('|')[-1]:
            cmds.select(children,add=1)
    cmds.delete()

    #选择组
    cmds.select(cmds.listRelatives(parent_group, children=True, fullPath=True)[0])
    
    #删除名称空间
    cmds.namespace(removeNamespace = obj_namespace, mergeNamespaceWithRoot = True)
    #获取最新对象
    obj_new = cmds.ls(sl=1)

    
    
    return obj_new,file_name


    
def exportFBX(fbx_export_path,file_name,select_obj,start_key,end_key):    
    
    #获取文件路径并创建新路径
    dir_path = fbx_export_path
    sc_folder = str(cmds.file(q=1,sn=1)).rsplit('_',1)[0].rsplit('/',1)[-1]
    sc_name = sc_folder.split('_')[1]
    text_path_old=dir_path+'/'+sc_name+'/'+sc_folder
    save_file_name='/'+str(cmds.file(q=1,sn=1)).rsplit('_',1)[0].rsplit('/',1)[-1]+'_an_'+file_name#创建文件名
    export_path=text_path_old+save_file_name
    print(export_path)
    #导出fbx
    cmds.select(select_obj)
    cmds.FBXResetExport()
    mel.eval('FBXExportFileVersion -v FBX201300')
    mel.eval('FBXExportSmoothingGroups -v true')
    mel.eval('FBXExportConstraints  -v false')
    mel.eval('FBXExportSkeletonDefinitions   -v false')
    mel.eval('FBXExportInputConnections -v false')
    mel.eval('FBXExportUpAxis y')
    mel.eval('FBXExportSmoothMesh -v true')
    #mel.eval('FBXExportSplitAnimationIntoTakes -v \"Take_001 " %s %s'%(start,end+start_offset+end_offset))
    #mel.eval('FBXExportSplitAnimationIntoTakes -c ')
    
    
    
    #创建对应FBX文件夹
    if not os.path.exists(export_path.rsplit('/',1)[0]):
        os.makedirs(export_path.rsplit('/',1)[0])
    
    if os.path.exists(export_path+'.fbx'):
        os.remove(export_path+'.fbx')
        print('delete '+export_path+'.fbx')

    
    #设置时间轴范围    
    cmds.playbackOptions( minTime=start_key, maxTime=end_key )
    #导出
    mel.eval('FBXExport -f "%s" -s'%(export_path))
    
    #eval('file -force -options "" -typ "FBX export" -pr -es "%s.fbx"'%(export_path))
    
    #ExportFbx(export_path)
    
    #创建收纳组
    try:
        cmds.select(select_obj)
        cmds.pickWalk( direction='up' )
        cmds.pickWalk( direction='up' )
        cmds.pickWalk( direction='up' )
        cmds.pickWalk( direction='up' )
        cmds.pickWalk( direction='up' )
        cmds.rename('Group_'+file_name.split(':',1)[0])
    except:
        pass   

    return export_path+'.fbx'          


    
    
# import maya.standalone as standalone
# standalone.initialize()



execute()
# cmds.file(new=True, force=True)
# cmds.unloadPlugin("mtoa", force=True)





